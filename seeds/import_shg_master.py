"""
Import TRESP SHG Members master data (16-field sheet) into MongoDB.

 Goal (as per your instruction):
- Insert EVERY Excel row into MongoDB (1 row = 1 document) so the final count becomes EXACTLY 140322.
- Do NOT skip blank Member Code rows.
- Do NOT collapse duplicates via (SHG Code, Member Code) upsert.
- Still build/maintain shg_master (derived) and geo masters (states/districts/blocks).

How this version works:
1) shg_members_master
   - Inserts ONE document per Excel row.
   - Adds a stable unique key per row: _import_key
     (based on row_number + SHG Code + Member Code; row_number guarantees uniqueness)
   - Uses UPSERT on _import_key so re-running does not duplicate rows.

2) shg_master (derived)
   - Aggregates unique Member Codes per SHG Code (blank member codes excluded from "Active Members" count)
   - Upserts SHG master by (State, District, Block, Gram Panchayat, Village, SHG Code)

Usage:
  python seeds/import_shg_master.py "seeds/data/TRESP_area_SHG_Members_Details.xlsx"

Env:
  Uses MONGO_URI and MONGO_DB from .env
"""

import os
import sys
import re
from datetime import datetime, timezone
from typing import Optional, Dict, Any, Tuple, Set

from dotenv import load_dotenv
from pymongo import MongoClient, UpdateOne
from openpyxl import load_workbook


EXPECTED_HEADERS = [
    "State", "District", "Block", "Gram Panchayat", "Village",
    "SHG Code", "SHG Name",
    "Member Code", "Member Name",
    "Designation in SHG", "Social Category", "Religion", "Education",
    "Disability", "Is head of Family", "Father/Mother/Spouse Name"
]


def _norm_header(s: str) -> str:
    s = str(s or "").strip().lower()
    s = re.sub(r"\s+", " ", s)
    return s


def _as_str(v):
    if v is None:
        return None
    s = str(v).strip()
    s = re.sub(r"\s+", " ", s).strip()
    return s if s else None


def _clean_name(v: Optional[str]) -> Optional[str]:
    if not v:
        return None
    s = re.sub(r"\s+", " ", str(v).strip())
    return s if s else None


def _mk_state_code(name: str) -> str:
    n = re.sub(r"\s+", " ", (name or "").strip()).upper()
    parts = [p for p in re.split(r"[^A-Z0-9]+", n) if p]
    if len(parts) >= 2:
        base = (parts[0][:1] + parts[1][:1])[:2]
    elif len(parts) == 1:
        base = parts[0][:2]
    else:
        base = "ST"
    base = re.sub(r"[^A-Z0-9]", "", base) or "ST"
    return base[:2]


def main(xlsx_path: str):
    if not os.path.exists(xlsx_path):
        raise FileNotFoundError(f"XLSX not found: {xlsx_path}")

    load_dotenv()

    mongo_uri = os.getenv("MONGO_URI")
    mongo_db = os.getenv("MONGO_DB", "pg_mis")

    if not mongo_uri:
        raise RuntimeError("MONGO_URI is not set in environment / .env")

    client = MongoClient(mongo_uri)
    db = client[mongo_db]

    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb.active

    # Header row (assumed row 1)
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    headers_norm = [_norm_header(h) for h in header_row]

    expected_norm = [_norm_header(h) for h in EXPECTED_HEADERS]

    # Map normalized header -> index
    idx: Dict[str, int] = {}
    missing = []
    for h in expected_norm:
        if h in headers_norm:
            idx[h] = headers_norm.index(h)
        else:
            missing.append(h)

    if missing:
        raise RuntimeError(
            "Missing required columns in XLSX: " + ", ".join(missing) +
            "\nFound columns: " + ", ".join([str(h) for h in headers_norm if h])
        )

    now_utc = lambda: datetime.now(timezone.utc)

    # ===============================
    # IMPORTANT: We want 1 Excel row = 1 doc
    # Use _import_key as the unique id for each row.
    # ===============================
    # Create index (safe if exists)
    try:
        db.shg_members_master.create_index([("_import_key", 1)], unique=True)
    except Exception:
        pass

    member_ops = []
    shg_agg: Dict[str, Dict[str, Any]] = {}

    geo_states: Set[str] = set()
    geo_districts: Set[Tuple[str, str]] = set()
    geo_blocks: Set[Tuple[str, str, str]] = set()

    batch_size = 1000
    total_rows = 0
    inserted_rows = 0

    # Iterate data rows
    for excel_row_number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        total_rows += 1
        if total_rows % 5000 == 0:
            print(f"Processed rows: {total_rows}")

        state = _clean_name(_as_str(row[idx["state"]]))
        district = _clean_name(_as_str(row[idx["district"]]))
        block = _clean_name(_as_str(row[idx["block"]]))
        gp = _clean_name(_as_str(row[idx["gram panchayat"]]))
        village = _clean_name(_as_str(row[idx["village"]]))

        shg_code = _as_str(row[idx["shg code"]])
        shg_name = _clean_name(_as_str(row[idx["shg name"]]))

        member_code = _as_str(row[idx["member code"]])  # may be None/blank
        member_name = _clean_name(_as_str(row[idx["member name"]]))

       

        # Track geo hierarchy
        if state:
            geo_states.add(state)
        if state and district:
            geo_districts.add((state, district))
        if state and district and block:
            geo_blocks.add((state, district, block))

        #  Unique per row import key (guarantees exact row count in DB)
        # row_number alone guarantees uniqueness, but we add shg/member for readability.
        import_key = f"TRESP|row={excel_row_number}|shg={shg_code or 'BLANK'}|mem={member_code or 'BLANK'}"

        mdoc = {
            "_import_key": import_key,  # <-- ensures 1 row = 1 document
            "State": state,
            "District": district,
            "Block": block,
            "Gram Panchayat": gp,
            "Village": village,
            "SHG Code": shg_code,
            "SHG Name": shg_name,
            "Member Code": member_code,  # may be None
            "Member Name": member_name,
            "Designation in SHG": _as_str(row[idx["designation in shg"]]),
            "Social Category": _as_str(row[idx["social category"]]),
            "Religion": _as_str(row[idx["religion"]]),
            "Education": _as_str(row[idx["education"]]),
            "Disability": _as_str(row[idx["disability"]]),
            "Is head of Family": _as_str(row[idx["is head of family"]]),
            "Father/Mother/Spouse Name": _as_str(row[idx["father/mother/spouse name"]]),
            "updated_at": now_utc(),
            "source": "TRESP",
        }

        #  Upsert by _import_key (NOT by SHG+Member)
        member_ops.append(
            UpdateOne(
                {"_import_key": import_key},
                {"$set": mdoc},
                upsert=True
            )
        )
        inserted_rows += 1

        # Aggregate SHG (for shg_master)
        if shg_code:
            agg = shg_agg.get(shg_code)
            if not agg:
                agg = {
                    "State": state,
                    "District": district,
                    "Block": block,
                    "Gram Panchayat": gp,
                    "Village": village,
                    "SHG Code": shg_code,
                    "SHG Name": shg_name,
                    "_members": set(),  # unique member_code set
                }
                shg_agg[shg_code] = agg

        # For Active Members count, only count non-empty member_code
        if member_code:
            agg["_members"].add(member_code)

        # Flush batch
        if len(member_ops) >= batch_size:
            db.shg_members_master.bulk_write(member_ops, ordered=False)
            member_ops = []

    # Flush remaining members
    if member_ops:
        db.shg_members_master.bulk_write(member_ops, ordered=False)

    # ===============================
    # Sync States / Districts / Blocks master
    # ===============================
    state_id_map: Dict[str, Any] = {}

    try:
        db.states.create_index([("code", 1)], unique=True)
    except Exception:
        pass

    for st in sorted(geo_states, key=lambda x: str(x).lower()):
        st_clean = _clean_name(st)
        if not st_clean:
            continue

        existing_by_name = db.states.find_one({"name": st_clean})
        if existing_by_name:
            state_id_map[st_clean] = existing_by_name["_id"]
            continue

        base = _mk_state_code(st_clean)
        code = base
        suffix = 0
        while db.states.find_one({"code": code}):
            suffix += 1
            code = f"{base}{suffix}"

        inserted_id = db.states.insert_one({
            "code": code,
            "name": st_clean,
            "created_at": now_utc(),
        }).inserted_id
        state_id_map[st_clean] = inserted_id

    district_id_map: Dict[Tuple[str, str], Any] = {}
    for st, dist in sorted(geo_districts, key=lambda x: (str(x[0]).lower(), str(x[1]).lower())):
        st_clean = _clean_name(st)
        dist_clean = _clean_name(dist)
        sid = state_id_map.get(st_clean)
        if not sid or not dist_clean:
            continue

        existing = db.districts.find_one({"name": dist_clean, "state_id": sid})
        if existing:
            district_id_map[(st_clean, dist_clean)] = existing["_id"]
        else:
            district_id_map[(st_clean, dist_clean)] = db.districts.insert_one({
                "name": dist_clean,
                "state_id": sid,
                "created_at": now_utc(),
            }).inserted_id

    for st, dist, blk in sorted(
        geo_blocks,
        key=lambda x: (str(x[0]).lower(), str(x[1]).lower(), str(x[2]).lower())
    ):
        st_clean = _clean_name(st)
        dist_clean = _clean_name(dist)
        blk_clean = _clean_name(blk)

        did = district_id_map.get((st_clean, dist_clean))
        if not did or not blk_clean:
            continue

        if not db.blocks.find_one({"name": blk_clean, "district_id": did}):
            db.blocks.insert_one({
                "name": blk_clean,
                "district_id": did,
                "created_at": now_utc(),
            })

    # ===============================
    # Upsert SHG master derived
    # ===============================
    shg_ops = []
    for shg_code, agg in shg_agg.items():
        shg_doc = {
            "State": agg["State"],
            "District": agg["District"],
            "Block": agg["Block"],
            "Gram Panchayat": agg["Gram Panchayat"],
            "Village": agg["Village"],
            "SHG Code": agg["SHG Code"],
            "SHG Name": agg["SHG Name"],
            "Active Members": len(agg["_members"]),
            "shg_nic_code": None,
            "Status": "Active",
            "updated_at": now_utc(),
            "source": "TRESP",
        }
        shg_ops.append(
            UpdateOne(
                {
                    "State": shg_doc["State"],
                    "District": shg_doc["District"],
                    "Block": shg_doc["Block"],
                    "Gram Panchayat": shg_doc["Gram Panchayat"],
                    "Village": shg_doc["Village"],
                    "SHG Code": shg_doc["SHG Code"],
                },
                {"$set": shg_doc},
                upsert=True
            )
        )
        if len(shg_ops) >= batch_size:
            db.shg_master.bulk_write(shg_ops, ordered=False)
            shg_ops = []

    if shg_ops:
        db.shg_master.bulk_write(shg_ops, ordered=False)

    # Helpful indexes
    db.shg_master.create_index([("State", 1), ("District", 1), ("Block", 1), ("Gram Panchayat", 1), ("Village", 1)])
    db.shg_master.create_index([("SHG Code", 1)])
    db.shg_members_master.create_index([("SHG Code", 1)])
    db.shg_members_master.create_index([("Member Code", 1)])

    print(f"Imported TRESP master into db={mongo_db}")
    print(f"Excel rows scanned (min_row=2): {total_rows}")
    print(f"Rows inserted/upserted into shg_members_master (with _import_key): {inserted_rows}")
    print(f"Derived SHGs in shg_master: {len(shg_agg)}")
    print("Collections: shg_master (derived), shg_members_master (raw rows, 1 row = 1 doc)")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print('Usage: python seeds/import_shg_master.py "seeds/data/TRESP_area_SHG_Members_Details.xlsx"')
        sys.exit(1)
    main(sys.argv[1])
